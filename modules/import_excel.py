"""
S.C.A.H. - Módulo de Importación desde Excel
Lee archivos .xlsx y .xls con formato específico de hoteles y huéspedes
"""

import os
import threading
import customtkinter as ctk
from tkinter import filedialog
from openpyxl import load_workbook
from datetime import datetime
from utils.geography import normalizar_huesped_geografia

try:
    import xlrd
    XLRD_DISPONIBLE = True
except ImportError:
    XLRD_DISPONIBLE = False


# ---- Adaptadores para que xlrd use la misma interfaz que openpyxl ----

class _XlrdCelda:
    """Envuelve un valor de celda xlrd para imitar cell.value de openpyxl."""
    def __init__(self, valor):
        self.value = valor


class _XlrdHoja:
    """Envuelve una hoja xlrd para imitar la interfaz de openpyxl."""
    def __init__(self, sheet, book):
        self._sheet = sheet
        self._book = book
        self.title = sheet.name

    def __getitem__(self, ref: str):
        """Accede a celdas con notación 'A1', 'B5', etc."""
        import re
        m = re.match(r'^([A-Za-z]+)(\d+)$', ref)
        if not m:
            return _XlrdCelda(None)
        col_letters = m.group(1).upper()
        row_num = int(m.group(2))
        # Convertir letras de columna a índice 0-based
        col_idx = 0
        for ch in col_letters:
            col_idx = col_idx * 26 + (ord(ch) - ord('A') + 1)
        col_idx -= 1  # 0-based
        row_idx = row_num - 1  # 0-based
        if row_idx < 0 or row_idx >= self._sheet.nrows:
            return _XlrdCelda(None)
        if col_idx < 0 or col_idx >= self._sheet.ncols:
            return _XlrdCelda(None)
        cell = self._sheet.cell(row_idx, col_idx)
        valor = cell.value
        # Convertir fechas de xlrd a datetime
        if cell.ctype == xlrd.XL_CELL_DATE and valor:
            try:
                dt_tuple = xlrd.xldate_as_tuple(valor, self._book.datemode)
                from datetime import datetime as _dt
                valor = _dt(*dt_tuple)
            except Exception:
                pass
        # Cadenas vacías → None (coherencia con openpyxl)
        if isinstance(valor, str) and valor.strip() == '':
            valor = None
        return _XlrdCelda(valor)


class _XlrdLibro:
    """Envuelve un workbook xlrd para imitar la interfaz de openpyxl."""
    def __init__(self, filepath):
        self._book = xlrd.open_workbook(filepath)
        self.sheetnames = self._book.sheet_names()

    def __getitem__(self, nombre):
        sheet = self._book.sheet_by_name(nombre)
        return _XlrdHoja(sheet, self._book)

    def close(self):
        self._book.release_resources()


def _abrir_workbook(filepath: str):
    """Abre un archivo Excel (.xlsx o .xls) y devuelve un objeto con interfaz unificada."""
    ext = os.path.splitext(filepath)[1].lower()
    if ext == '.xls':
        if not XLRD_DISPONIBLE:
            raise ImportError(
                "Se necesita la librería 'xlrd' para leer archivos .xls.\n"
                "Instálala con: pip install xlrd"
            )
        return _XlrdLibro(filepath)
    else:
        return load_workbook(filepath, read_only=True, data_only=True)

from config import EXCEL_HOTEL_MAP, EXCEL_HUESPED_COLS, EXCEL_HUESPED_START_ROW
from database.connection import db
from ui.themes import COLORS, obtener_fuente
from ui.components import TablaScrollable, BarraProgreso
from ui.dialogs import mostrar_exito, mostrar_error, mostrar_advertencia, confirmar
from utils.validators import (validar_dni, validar_fecha, validar_texto_obligatorio,
                               validar_edad, sanitizar_texto)
from utils.formatters import formato_fecha
from utils.logger import log_info, log_error, Auditoria


class ImportExcelModule(ctk.CTkFrame):
    """Módulo para importar datos desde archivos Excel."""

    def __init__(self, parent, usuario: dict):
        super().__init__(parent, fg_color="transparent")

        self.usuario = usuario
        self.archivos_seleccionados = []
        self.datos_preview = []
        self.datos_hotel_preview = []

        self._crear_ui()

    def _crear_ui(self):
        """Crea la interfaz del módulo de importación."""
        # --- Sección superior: selección de archivos ---
        select_frame = ctk.CTkFrame(self, fg_color=COLORS["bg_card"], corner_radius=10)
        select_frame.pack(fill="x", pady=(0, 10))

        header = ctk.CTkFrame(select_frame, fg_color="transparent")
        header.pack(fill="x", padx=20, pady=(15, 10))

        ctk.CTkLabel(
            header, text="Seleccionar archivos Excel",
            font=obtener_fuente("heading"),
            anchor="w"
        ).pack(side="left")

        btn_frame = ctk.CTkFrame(header, fg_color="transparent")
        btn_frame.pack(side="right")

        ctk.CTkButton(
            btn_frame, text="📂 Seleccionar archivos",
            font=obtener_fuente("button"),
            height=38,
            command=self._seleccionar_archivos
        ).pack(side="left", padx=5)

        ctk.CTkButton(
            btn_frame, text="🗂️ Seleccionar carpeta",
            font=obtener_fuente("button"),
            height=38,
            fg_color=COLORS["secondary"],
            hover_color=COLORS["secondary_hover"],
            command=self._seleccionar_carpeta
        ).pack(side="left", padx=5)

        # Lista de archivos seleccionados
        self.label_archivos = ctk.CTkLabel(
            select_frame,
            text="No se han seleccionado archivos",
            font=obtener_fuente("small"),
            text_color=COLORS["text_secondary"],
            anchor="w",
            wraplength=900
        )
        self.label_archivos.pack(fill="x", padx=20, pady=(0, 15))

        # --- Sección media: vista previa ---
        preview_frame = ctk.CTkFrame(self, fg_color=COLORS["bg_card"], corner_radius=10)
        preview_frame.pack(fill="both", expand=True, pady=(0, 10))

        prev_header = ctk.CTkFrame(preview_frame, fg_color="transparent")
        prev_header.pack(fill="x", padx=20, pady=(15, 10))

        ctk.CTkLabel(
            prev_header, text="Vista previa de datos",
            font=obtener_fuente("heading"),
            anchor="w"
        ).pack(side="left")

        self.label_conteo = ctk.CTkLabel(
            prev_header, text="",
            font=obtener_fuente("small"),
            text_color=COLORS["text_secondary"]
        )
        self.label_conteo.pack(side="right")

        # Pestañas: datos hotel / datos huéspedes
        tab_frame = ctk.CTkFrame(preview_frame, fg_color="transparent")
        tab_frame.pack(fill="x", padx=20, pady=(0, 5))

        self.btn_tab_hotel = ctk.CTkButton(
            tab_frame, text="🏨 Hoteles detectados", width=180, height=32,
            font=obtener_fuente("small_bold"),
            fg_color=COLORS["primary"],
            command=lambda: self._cambiar_tab("hotel")
        )
        self.btn_tab_hotel.pack(side="left", padx=(0, 5))

        self.btn_tab_huesped = ctk.CTkButton(
            tab_frame, text="👤 Huéspedes detectados", width=180, height=32,
            font=obtener_fuente("small_bold"),
            fg_color=COLORS["sidebar_hover"],
            command=lambda: self._cambiar_tab("huesped")
        )
        self.btn_tab_huesped.pack(side="left")

        # Tabla de hoteles
        self.tabla_hotel = TablaScrollable(
            preview_frame,
            columnas=[
                {"id": "archivo", "texto": "Archivo/Hoja", "ancho": 200},
                {"id": "nombre", "texto": "Hotel", "ancho": 200},
                {"id": "nro_orden", "texto": "Nro. Orden", "ancho": 100},
                {"id": "direccion", "texto": "Dirección", "ancho": 250},
                {"id": "ciudad", "texto": "Ciudad/Localidad", "ancho": 150},
            ],
            altura=6
        )

        # Tabla de huéspedes
        self.tabla_huesped = TablaScrollable(
            preview_frame,
            columnas=[
                {"id": "hotel", "texto": "Hotel", "ancho": 150},
                {"id": "nacionalidad", "texto": "Nacionalidad", "ancho": 100},
                {"id": "procedencia", "texto": "Procedencia", "ancho": 120},
                {"id": "apellido_nombre", "texto": "Apellido y Nombre", "ancho": 180},
                {"id": "dni_pasaporte", "texto": "DNI/Pasaporte", "ancho": 110},
                {"id": "fecha_nacimiento", "texto": "Fec. Nacimiento", "ancho": 110},
                {"id": "edad", "texto": "Edad", "ancho": 60},
                {"id": "profesion", "texto": "Profesión", "ancho": 120},
                {"id": "fecha_entrada", "texto": "Entrada", "ancho": 100},
                {"id": "fecha_salida", "texto": "Salida", "ancho": 100},
            ],
            altura=10
        )

        # Mostrar tabla de huéspedes por defecto
        self.tabla_huesped.pack(fill="both", expand=True, padx=20, pady=(0, 15))
        self.tab_actual = "huesped"

        # --- Sección inferior: barra de progreso + botones ---
        bottom_frame = ctk.CTkFrame(self, fg_color=COLORS["bg_card"], corner_radius=10)
        bottom_frame.pack(fill="x")

        self.barra_progreso = BarraProgreso(bottom_frame, "Esperando datos...")
        self.barra_progreso.pack(fill="x", padx=20, pady=(15, 10))

        # Info de errores
        self.label_errores = ctk.CTkLabel(
            bottom_frame, text="",
            font=obtener_fuente("small"),
            text_color=COLORS["warning"],
            anchor="w",
            wraplength=700
        )
        self.label_errores.pack(fill="x", padx=20, pady=(0, 5))

        # Botones de acción
        action_frame = ctk.CTkFrame(bottom_frame, fg_color="transparent")
        action_frame.pack(fill="x", padx=20, pady=(0, 15))

        self.btn_importar = ctk.CTkButton(
            action_frame, text="✅ Importar datos",
            font=obtener_fuente("button"),
            height=42,
            fg_color=COLORS["success"],
            hover_color="#388E3C",
            state="disabled",
            command=self._importar_datos
        )
        self.btn_importar.pack(side="right", padx=5)

        self.btn_limpiar = ctk.CTkButton(
            action_frame, text="🗑️ Limpiar",
            font=obtener_fuente("button"),
            height=42,
            fg_color=COLORS["error"],
            hover_color="#C62828",
            command=self._limpiar
        )
        self.btn_limpiar.pack(side="right", padx=5)

    def _cambiar_tab(self, tab: str):
        """Cambia entre las pestañas de hoteles y huéspedes."""
        if tab == "hotel":
            self.tabla_huesped.pack_forget()
            self.tabla_hotel.pack(fill="both", expand=True, padx=20, pady=(0, 15))
            self.btn_tab_hotel.configure(fg_color=COLORS["primary"])
            self.btn_tab_huesped.configure(fg_color=COLORS["sidebar_hover"])
        else:
            self.tabla_hotel.pack_forget()
            self.tabla_huesped.pack(fill="both", expand=True, padx=20, pady=(0, 15))
            self.btn_tab_huesped.configure(fg_color=COLORS["primary"])
            self.btn_tab_hotel.configure(fg_color=COLORS["sidebar_hover"])
        self.tab_actual = tab

    def _seleccionar_archivos(self):
        """Abre diálogo para seleccionar archivos Excel."""
        archivos = filedialog.askopenfilenames(
            title="Seleccionar archivos Excel",
            filetypes=[("Archivos Excel", "*.xlsx *.xls"), ("Todos", "*.*")]
        )
        if archivos:
            self.archivos_seleccionados = list(archivos)
            self._mostrar_archivos()
            self._procesar_archivos()

    def _seleccionar_carpeta(self):
        """Abre diálogo para seleccionar una carpeta con archivos Excel."""
        carpeta = filedialog.askdirectory(title="Seleccionar carpeta con archivos Excel")
        if carpeta:
            archivos = []
            for f in os.listdir(carpeta):
                if f.lower().endswith(('.xlsx', '.xls')) and not f.startswith('~'):
                    archivos.append(os.path.join(carpeta, f))
            if archivos:
                self.archivos_seleccionados = archivos
                self._mostrar_archivos()
                self._procesar_archivos()
            else:
                mostrar_advertencia(self, "Sin archivos",
                                    "No se encontraron archivos Excel en la carpeta seleccionada.")

    def _mostrar_archivos(self):
        """Muestra los archivos seleccionados en la UI."""
        nombres = [os.path.basename(f) for f in self.archivos_seleccionados]
        if len(nombres) > 5:
            texto = ", ".join(nombres[:5]) + f" y {len(nombres) - 5} más..."
        else:
            texto = ", ".join(nombres)
        self.label_archivos.configure(
            text=f"📁 {len(nombres)} archivo(s): {texto}",
            text_color=COLORS["text_primary"]
        )

    def _procesar_archivos(self):
        """Procesa los archivos Excel seleccionados en un hilo secundario."""
        self.datos_preview = []
        self.datos_hotel_preview = []

        self.barra_progreso.actualizar(0, "Procesando archivos...")
        # Deshabilitar botones mientras se procesa
        self._set_botones_estado("disabled")

        archivos = list(self.archivos_seleccionados)

        def tarea():
            datos_preview = []
            datos_hotel_preview = []
            errores = []
            total_hojas = 0

            for idx, archivo in enumerate(archivos):
                try:
                    wb = _abrir_workbook(archivo)
                    nombre_archivo = os.path.basename(archivo)

                    for hoja_nombre in wb.sheetnames:
                        hoja = wb[hoja_nombre]
                        total_hojas += 1

                        hotel_data = self._leer_datos_hotel(hoja, nombre_archivo, hoja_nombre)
                        if hotel_data:
                            datos_hotel_preview.append(hotel_data)

                        huespedes = self._leer_huespedes(hoja, hotel_data)
                        datos_preview.extend(huespedes)

                    wb.close()

                except Exception as e:
                    error_msg = f"Error en '{os.path.basename(archivo)}': {str(e)}"
                    errores.append(error_msg)
                    log_error(error_msg, e)

                progreso = (idx + 1) / len(archivos)
                self.after(0, lambda p=progreso, i=idx: self.barra_progreso.actualizar(
                    p, f"Procesando {i + 1}/{len(archivos)} archivos..."
                ))

            # Actualizar UI en el hilo principal
            self.after(0, lambda: self._procesar_archivos_completado(
                datos_preview, datos_hotel_preview, errores, total_hojas
            ))

        threading.Thread(target=tarea, daemon=True).start()

    def _set_botones_estado(self, estado: str):
        """Habilita o deshabilita los botones de acción."""
        try:
            for widget in self.winfo_children():
                self._set_botones_recursivo(widget, estado)
        except Exception:
            pass

    def _set_botones_recursivo(self, widget, estado: str):
        """Recursivamente establece el estado de botones CTkButton."""
        try:
            if isinstance(widget, ctk.CTkButton):
                widget.configure(state=estado)
            for child in widget.winfo_children():
                self._set_botones_recursivo(child, estado)
        except Exception:
            pass

    def _procesar_archivos_completado(self, datos_preview, datos_hotel_preview, errores, total_hojas):
        """Callback en el hilo principal con los resultados del procesamiento."""
        self.datos_preview = datos_preview
        self.datos_hotel_preview = datos_hotel_preview

        self.barra_progreso.completar(
            f"Procesado: {len(self.archivos_seleccionados)} archivos, "
            f"{total_hojas} hojas, {len(self.datos_preview)} huéspedes encontrados"
        )

        self.tabla_hotel.cargar_datos([
            (h["archivo_hoja"], h["nombre"], h["nro_orden"],
             h["direccion"], h["ciudad"])
            for h in self.datos_hotel_preview
        ])

        self.tabla_huesped.cargar_datos([
            (h["hotel_nombre"], h["nacionalidad"], h["procedencia"],
             h["apellido_nombre"], h["dni_pasaporte"],
             formato_fecha(h.get("fecha_nacimiento")),
             str(h.get("edad", "")),
             h["profesion"],
             formato_fecha(h.get("fecha_entrada")),
             formato_fecha(h.get("fecha_salida")))
            for h in self.datos_preview
        ])

        self.label_conteo.configure(
            text=f"{len(self.datos_hotel_preview)} hoteles | {len(self.datos_preview)} huéspedes"
        )

        if errores:
            self.label_errores.configure(text="⚠ " + " | ".join(errores[:3]))
        else:
            self.label_errores.configure(text="")

        # Rehabilitar botones
        self._set_botones_estado("normal")
        if self.datos_preview:
            self.btn_importar.configure(state="normal")
        else:
            self.btn_importar.configure(state="disabled")

    def _leer_datos_hotel(self, hoja, nombre_archivo: str, hoja_nombre: str) -> dict:
        """Lee los datos del hotel de una hoja del Excel.
        Intenta las posiciones configuradas y también busca en las primeras filas."""
        try:
            # Intentar posiciones configuradas primero
            nombre = self._valor_celda(hoja, EXCEL_HOTEL_MAP["nombre"]["col"], EXCEL_HOTEL_MAP["nombre"]["row"])
            nro_orden = self._valor_celda(hoja, EXCEL_HOTEL_MAP["nro_orden"]["col"], EXCEL_HOTEL_MAP["nro_orden"]["row"])
            direccion = self._valor_celda(hoja, EXCEL_HOTEL_MAP["direccion"]["col"], EXCEL_HOTEL_MAP["direccion"]["row"])
            ciudad = self._valor_celda(hoja, EXCEL_HOTEL_MAP["ciudad_localidad"]["col"], EXCEL_HOTEL_MAP["ciudad_localidad"]["row"])

            # Si no encontró nombre en la posición configurada, buscar en las primeras filas de la columna A
            if not nombre or str(nombre).strip() == "":
                for r in range(1, 10):
                    val = self._valor_celda(hoja, "A", r)
                    if val and str(val).strip() and len(str(val).strip()) > 2:
                        texto = str(val).strip().lower()
                        # Saltar encabezados genéricos
                        if texto not in ("hotel", "nombre", "nro", "orden", "dirección",
                                         "direccion", "ciudad", "localidad"):
                            nombre = val
                            break

            # También intentar usar el nombre de la hoja como nombre de hotel si no hay datos
            if not nombre or str(nombre).strip() == "":
                nombre = hoja_nombre

            return {
                "archivo_hoja": f"{nombre_archivo} / {hoja_nombre}",
                "nombre": sanitizar_texto(str(nombre)) if nombre else f"Hotel ({hoja_nombre})",
                "nro_orden": sanitizar_texto(str(nro_orden)) if nro_orden else "",
                "direccion": sanitizar_texto(str(direccion)) if direccion else "",
                "ciudad": sanitizar_texto(str(ciudad)) if ciudad else "",
                "hoja": hoja_nombre
            }
        except Exception as e:
            log_error(f"Error leyendo datos de hotel en {hoja_nombre}", e)
            return {
                "archivo_hoja": f"{nombre_archivo} / {hoja_nombre}",
                "nombre": f"Hotel ({hoja_nombre})",
                "nro_orden": "", "direccion": "", "ciudad": "",
                "hoja": hoja_nombre
            }

    def _detectar_fila_inicio(self, hoja) -> int:
        """Detecta automáticamente la fila donde comienzan los datos de huéspedes.
        Escanea desde la fila 1 buscando la primera fila que tenga datos válidos
        en las columnas de huéspedes (no encabezados)."""
        columnas_a_verificar = [
            EXCEL_HUESPED_COLS["apellido_nombre"],
            EXCEL_HUESPED_COLS["dni_pasaporte"],
            EXCEL_HUESPED_COLS["nacionalidad"],
        ]
        # Palabras clave de encabezados que NO son datos
        encabezados = {
            "apellido", "nombre", "apellido y nombre", "apellido_nombre",
            "dni", "pasaporte", "dni/pasaporte", "documento",
            "nacionalidad", "procedencia", "profesion", "profesión",
            "edad", "entrada", "salida", "nacimiento",
            "fecha", "hotel", "nro", "orden", "dirección", "direccion",
            "ciudad", "localidad", "fec. nacimiento", "fec. entrada",
            "fec. salida", "fecha nacimiento", "fecha entrada", "fecha salida",
        }

        for fila in range(1, 50):  # Escanear hasta fila 50
            for col in columnas_a_verificar:
                valor = self._valor_celda(hoja, col, fila)
                if valor is not None:
                    texto = str(valor).strip().lower()
                    if texto and texto not in encabezados and len(texto) > 1:
                        # Encontramos datos reales, no un encabezado
                        return fila
        return EXCEL_HUESPED_START_ROW  # Fallback al valor configurado

    def _fila_tiene_datos(self, hoja, fila: int) -> bool:
        """Verifica si una fila tiene datos en al menos una columna de huéspedes."""
        for col in EXCEL_HUESPED_COLS.values():
            valor = self._valor_celda(hoja, col, fila)
            if valor is not None and str(valor).strip() != "":
                return True
        return False

    def _obtener_max_fila(self, hoja) -> int:
        """Obtiene la última fila con datos en la hoja."""
        # Para openpyxl
        if hasattr(hoja, 'max_row') and hoja.max_row:
            return hoja.max_row
        # Para xlrd adapter
        if hasattr(hoja, '_sheet') and hasattr(hoja._sheet, 'nrows'):
            return hoja._sheet.nrows
        return 10000  # Fallback

    def _leer_huespedes(self, hoja, hotel_data: dict) -> list:
        """Lee los datos de huéspedes de una hoja del Excel.
        Detecta automáticamente la fila de inicio y escanea todas las filas."""
        huespedes = []

        # Detectar fila de inicio automáticamente
        fila = self._detectar_fila_inicio(hoja)
        max_fila = self._obtener_max_fila(hoja)
        filas_vacias_consecutivas = 0
        MAX_FILAS_VACIAS = 15  # Tolerar hasta 15 filas vacías antes de parar

        while fila <= max_fila + MAX_FILAS_VACIAS:
            # Verificar si la fila tiene algún dato en cualquier columna de huéspedes
            if not self._fila_tiene_datos(hoja, fila):
                filas_vacias_consecutivas += 1
                if filas_vacias_consecutivas >= MAX_FILAS_VACIAS:
                    break
                fila += 1
                continue

            filas_vacias_consecutivas = 0

            # Leer apellido/nombre - puede estar vacío si hay otros datos
            apellido = self._valor_celda(hoja, EXCEL_HUESPED_COLS["apellido_nombre"], fila)
            if not apellido or str(apellido).strip() == "":
                # La fila tiene datos en otras columnas pero no nombre - intentar usar DNI como referencia
                dni_val = self._valor_celda(hoja, EXCEL_HUESPED_COLS["dni_pasaporte"], fila)
                if not dni_val or str(dni_val).strip() == "":
                    fila += 1
                    continue

            try:
                # Leer fecha de nacimiento
                fn_raw = self._valor_celda(hoja, EXCEL_HUESPED_COLS["fecha_nacimiento"], fila)
                _, _, fecha_nac = validar_fecha(fn_raw, permite_vacio=True)

                # Leer fecha de entrada
                fe_raw = self._valor_celda(hoja, EXCEL_HUESPED_COLS["fecha_entrada"], fila)
                _, _, fecha_entrada = validar_fecha(fe_raw, permite_vacio=True)

                # Leer fecha de salida
                fs_raw = self._valor_celda(hoja, EXCEL_HUESPED_COLS["fecha_salida"], fila)
                _, _, fecha_salida = validar_fecha(fs_raw, permite_vacio=True)

                # Leer edad
                edad_raw = self._valor_celda(hoja, EXCEL_HUESPED_COLS["edad"], fila)
                _, _, edad = validar_edad(edad_raw)

                huesped = {
                    "hotel_nombre": hotel_data.get("nombre", "") if hotel_data else "",
                    "hotel_data": hotel_data,
                    "nacionalidad": sanitizar_texto(
                        str(self._valor_celda(hoja, EXCEL_HUESPED_COLS["nacionalidad"], fila) or "")),
                    "procedencia": sanitizar_texto(
                        str(self._valor_celda(hoja, EXCEL_HUESPED_COLS["procedencia"], fila) or "")),
                    "apellido_nombre": sanitizar_texto(str(apellido)),
                    "dni_pasaporte": sanitizar_texto(
                        str(self._valor_celda(hoja, EXCEL_HUESPED_COLS["dni_pasaporte"], fila) or "")),
                    "fecha_nacimiento": fecha_nac,
                    "edad": edad,
                    "profesion": sanitizar_texto(
                        str(self._valor_celda(hoja, EXCEL_HUESPED_COLS["profesion"], fila) or "")),
                    "fecha_entrada": fecha_entrada,
                    "fecha_salida": fecha_salida,
                }
                huespedes.append(normalizar_huesped_geografia(huesped))

            except Exception as e:
                log_error(f"Error leyendo huésped en fila {fila}", e)

            fila += 1

        return huespedes

    def _valor_celda(self, hoja, columna: str, fila: int):
        """Obtiene el valor de una celda específica."""
        try:
            celda = hoja[f"{columna}{fila}"]
            valor = celda.value
            # Limpiar valores que sean solo espacios en blanco
            if isinstance(valor, str) and valor.strip() == "":
                return None
            return valor
        except (KeyError, IndexError, AttributeError):
            return None
        except Exception:
            return None

    def _importar_datos(self):
        """Importa los datos procesados a la base de datos en un hilo secundario."""
        if not self.datos_preview:
            mostrar_advertencia(self, "Sin datos", "No hay datos para importar.")
            return

        if not confirmar(self, "Confirmar importación",
                         f"Se importarán {len(self.datos_preview)} registros de huéspedes "
                         f"de {len(self.datos_hotel_preview)} hoteles.\n\n¿Desea continuar?"):
            return

        self.btn_importar.configure(state="disabled", text="Importando...")
        self.btn_limpiar.configure(state="disabled")
        self.barra_progreso.actualizar(0, "Importando datos a la base de datos...")

        datos = list(self.datos_preview)
        archivos = list(self.archivos_seleccionados)
        usuario_id = self.usuario["id"]

        def tarea():
            importados = 0
            errores = 0
            duplicados = 0
            error_general = None

            try:
                conn = db.obtener_conexion()
                if not conn:
                    self.after(0, lambda: mostrar_error(self, "Error de conexión",
                                  "No se pudo conectar a la base de datos."))
                    return

                cursor = conn.cursor()
                hoteles_cache = {}

                for i, huesped in enumerate(datos):
                    try:
                        hotel_data = huesped.get("hotel_data", {})
                        hotel_key = hotel_data.get("nombre", "")

                        if hotel_key not in hoteles_cache:
                            hotel_id = self._obtener_o_crear_hotel(cursor, hotel_data)
                            hoteles_cache[hotel_key] = hotel_id
                        else:
                            hotel_id = hoteles_cache[hotel_key]

                        if self._es_duplicado(cursor, hotel_id, huesped):
                            duplicados += 1
                            continue

                        cursor.execute("""
                            INSERT INTO huespedes (
                                hotel_id, nacionalidad, procedencia, apellido_nombre,
                                dni_pasaporte, fecha_nacimiento, edad, profesion,
                                fecha_entrada, fecha_salida, origen_carga, usuario_carga_id
                            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, 'excel', %s)
                        """, (
                            hotel_id,
                            huesped["nacionalidad"],
                            huesped["procedencia"],
                            huesped["apellido_nombre"],
                            huesped["dni_pasaporte"],
                            huesped["fecha_nacimiento"],
                            huesped["edad"],
                            huesped["profesion"],
                            huesped["fecha_entrada"],
                            huesped["fecha_salida"],
                            usuario_id
                        ))
                        importados += 1

                    except Exception as e:
                        errores += 1
                        log_error(f"Error importando huésped: {huesped.get('apellido_nombre', '?')}", e)

                    if (i + 1) % 10 == 0 or i == len(datos) - 1:
                        progreso = (i + 1) / len(datos)
                        self.after(0, lambda p=progreso, n=i+1: self.barra_progreso.actualizar(
                            p, f"Importando {n}/{len(datos)}..."
                        ))

                for archivo in archivos:
                    cursor.execute("""
                        INSERT INTO importaciones_log (
                            archivo_nombre, fecha_importacion, usuario_id,
                            registros_importados, registros_error, registros_duplicados, estado
                        ) VALUES (%s, %s, %s, %s, %s, %s, %s)
                    """, (
                        os.path.basename(archivo),
                        datetime.now(),
                        usuario_id,
                        importados,
                        errores,
                        duplicados,
                        "completado" if errores == 0 else "parcial"
                    ))

                conn.commit()
                cursor.close()
                db.liberar_conexion(conn)

                try:
                    conn_aud = db.obtener_conexion()
                    if conn_aud:
                        auditoria = Auditoria(conn_aud)
                        auditoria.registrar(
                            usuario_id,
                            "importar_excel",
                            "huespedes",
                            detalle=f"Importados: {importados}, Errores: {errores}, Duplicados: {duplicados}"
                        )
                        db.liberar_conexion(conn_aud)
                except Exception:
                    pass

            except Exception as e:
                log_error("Error general en importación", e)
                error_general = str(e)

            self.after(0, lambda: self._importar_completado(importados, errores, duplicados, error_general))

        threading.Thread(target=tarea, daemon=True).start()

    def _importar_completado(self, importados, errores, duplicados, error_general):
        """Callback en el hilo principal con los resultados de la importación."""
        self.btn_importar.configure(state="normal", text="✅ Importar datos")
        self.btn_limpiar.configure(state="normal")

        if error_general:
            mostrar_error(self, "Error de importación",
                          f"Ocurrió un error durante la importación:\n{error_general}")
            return

        self.barra_progreso.completar(
            f"Importación completada: {importados} importados, "
            f"{duplicados} duplicados, {errores} errores"
        )

        mensaje = (f"Importación completada:\n\n"
                   f"✅ Importados: {importados}\n"
                   f"⚠️ Duplicados omitidos: {duplicados}\n"
                   f"❌ Errores: {errores}")

        if errores == 0:
            mostrar_exito(self, "Importación exitosa", mensaje)
        else:
            mostrar_advertencia(self, "Importación parcial", mensaje)

        log_info(f"Importación Excel: {importados} importados, {duplicados} duplicados, {errores} errores")

    def _obtener_o_crear_hotel(self, cursor, hotel_data: dict) -> int:
        """Obtiene el ID de un hotel existente o lo crea.
        Si ya existe, actualiza campos vacíos con datos del Excel."""
        nombre = hotel_data.get("nombre", "Sin nombre")

        # Buscar hotel existente
        cursor.execute(
            "SELECT id FROM hoteles WHERE LOWER(nombre) = LOWER(%s)", (nombre,)
        )
        resultado = cursor.fetchone()

        if resultado:
            hotel_id = resultado[0]
            # Actualizar campos vacíos del hotel con datos del Excel
            updates = []
            params = []
            for campo_db, campo_data in [("direccion", "direccion"),
                                          ("ciudad_localidad", "ciudad"),
                                          ("nro_orden", "nro_orden")]:
                valor = hotel_data.get(campo_data, "")
                if valor:
                    updates.append(f"{campo_db} = COALESCE(NULLIF({campo_db}, ''), %s)")
                    params.append(valor)
            if updates:
                params.append(hotel_id)
                cursor.execute(
                    f"UPDATE hoteles SET {', '.join(updates)} WHERE id = %s",
                    tuple(params)
                )
            return hotel_id

        # Crear nuevo hotel
        cursor.execute("""
            INSERT INTO hoteles (nombre, nro_orden, direccion, ciudad_localidad, usuario_registro_id)
            VALUES (%s, %s, %s, %s, %s) RETURNING id
        """, (
            nombre,
            hotel_data.get("nro_orden", ""),
            hotel_data.get("direccion", ""),
            hotel_data.get("ciudad", ""),
            self.usuario["id"]
        ))
        return cursor.fetchone()[0]

    def _es_duplicado(self, cursor, hotel_id: int, huesped: dict) -> bool:
        """Verifica si un huésped ya existe (mismo DNI + hotel + fecha entrada)."""
        dni = huesped.get("dni_pasaporte", "")
        fecha_entrada = huesped.get("fecha_entrada")

        if not dni or not fecha_entrada:
            return False

        cursor.execute("""
            SELECT id FROM huespedes 
            WHERE hotel_id = %s AND dni_pasaporte = %s AND fecha_entrada = %s
        """, (hotel_id, dni, fecha_entrada))

        return cursor.fetchone() is not None

    def _limpiar(self):
        """Limpia todos los datos y resetea la interfaz."""
        self.archivos_seleccionados = []
        self.datos_preview = []
        self.datos_hotel_preview = []
        self.tabla_hotel.limpiar()
        self.tabla_huesped.limpiar()
        self.label_archivos.configure(
            text="No se han seleccionado archivos",
            text_color=COLORS["text_secondary"]
        )
        self.label_conteo.configure(text="")
        self.label_errores.configure(text="")
        self.barra_progreso.actualizar(0, "Esperando datos...")
        self.btn_importar.configure(state="disabled")
