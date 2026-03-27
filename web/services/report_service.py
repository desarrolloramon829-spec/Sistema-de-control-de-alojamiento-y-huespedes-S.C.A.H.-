"""
S.C.A.H. Web - Servicio de Reportes
Generación de reportes Excel y PDF server-side.
"""

import sys
import os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..'))

from datetime import datetime
from database.connection import db
from utils.logger import log_info, log_error
from utils.formatters import formato_nombre_archivo
from web.services.stats_service import obtener_segmentacion_geografica


def obtener_datos_reporte(tipo: str, params: dict = None) -> tuple[list, str]:
    """
    Obtiene datos para un tipo de reporte.
    Retorna (datos, titulo).
    """
    if tipo == "general":
        return _reporte_general()
    elif tipo == "por_hotel":
        return _reporte_por_hotel()
    elif tipo == "por_fechas":
        return _reporte_por_fechas(params.get("desde"), params.get("hasta"))
    elif tipo == "estadistico":
        return _reporte_estadistico()
    elif tipo == "geografico":
        return _reporte_geografico()
    elif tipo == "importaciones":
        return _reporte_importaciones()
    elif tipo == "auditoria":
        return _reporte_auditoria()
    return [], "Reporte no encontrado"


def _reporte_general():
    datos = db.ejecutar_query("""
        SELECT hu.apellido_nombre, hu.dni_pasaporte, hu.nacionalidad,
               hu.procedencia, hu.profesion, hu.edad,
               hu.fecha_entrada, hu.fecha_salida, h.nombre as hotel,
               h.ciudad_localidad, hu.habitacion, hu.domicilio,
               hu.destino, hu.movilidad, hu.telefono
        FROM huespedes hu
        LEFT JOIN hoteles h ON hu.hotel_id = h.id
        ORDER BY hu.apellido_nombre
    """, fetch=True)
    return [dict(d) for d in datos] if datos else [], "Listado General de Huéspedes"


def _reporte_por_hotel():
    datos = db.ejecutar_query("""
        SELECT h.nombre as hotel, h.ciudad_localidad,
               hu.apellido_nombre, hu.dni_pasaporte, hu.nacionalidad,
               hu.profesion, hu.edad,
               hu.fecha_entrada, hu.fecha_salida,
               hu.habitacion, hu.telefono
        FROM huespedes hu
        JOIN hoteles h ON hu.hotel_id = h.id
        ORDER BY h.nombre, hu.apellido_nombre
    """, fetch=True)
    return [dict(d) for d in datos] if datos else [], "Huéspedes por Hotel"


def _reporte_por_fechas(desde_str, hasta_str):
    try:
        desde = datetime.strptime(desde_str, "%Y-%m-%d").date()
        hasta = datetime.strptime(hasta_str, "%Y-%m-%d").date()
    except (ValueError, TypeError):
        return [], "Fechas inválidas"

    datos = db.ejecutar_query("""
        SELECT hu.apellido_nombre, hu.dni_pasaporte, hu.nacionalidad,
               hu.procedencia, hu.profesion, hu.edad,
               hu.fecha_entrada, hu.fecha_salida, h.nombre as hotel,
               hu.habitacion, hu.telefono
        FROM huespedes hu
        LEFT JOIN hoteles h ON hu.hotel_id = h.id
        WHERE hu.fecha_entrada BETWEEN %s AND %s
        ORDER BY hu.fecha_entrada, hu.apellido_nombre
    """, (desde, hasta), fetch=True)

    titulo = f"Huéspedes del {desde.strftime('%d/%m/%Y')} al {hasta.strftime('%d/%m/%Y')}"
    return [dict(d) for d in datos] if datos else [], titulo


def _reporte_estadistico():
    """Combina estadísticas de nacionalidad, profesión y procedencia."""
    datos_combinados = []

    for campo, nombre in [("nacionalidad", "Nacionalidad"),
                          ("profesion", "Profesión"),
                          ("procedencia", "Procedencia")]:
        resultados = db.ejecutar_query(f"""
            SELECT {campo} as concepto, COUNT(*) as total
            FROM huespedes
            WHERE {campo} IS NOT NULL AND {campo} != ''
            GROUP BY {campo} ORDER BY total DESC
        """, fetch=True)
        if resultados:
            for d in resultados[:10]:
                datos_combinados.append({
                    "categoria": nombre,
                    "concepto": d["concepto"],
                    "total": d["total"]
                })

    return datos_combinados, "Reporte Estadístico"


def _reporte_geografico():
    data = obtener_segmentacion_geografica()
    datos = []

    resumen = data.get('resumen', {})
    datos.extend([
        {
            'segmento': 'Resumen general',
            'agrupacion': 'Argentinos',
            'detalle': 'Total clasificado como argentino',
            'total': resumen.get('argentinos', 0),
        },
        {
            'segmento': 'Resumen general',
            'agrupacion': 'Extranjeros',
            'detalle': 'Total clasificado como extranjero',
            'total': resumen.get('extranjeros', 0),
        },
        {
            'segmento': 'Resumen general',
            'agrupacion': 'Sin clasificar',
            'detalle': 'Registros sin datos suficientes o no reconocidos',
            'total': resumen.get('sin_clasificar', 0),
        },
    ])

    for label, value in zip(data.get('series', {}).get('argentinos', {}).get('labels', []), data.get('series', {}).get('argentinos', {}).get('values', [])):
        datos.append({
            'segmento': 'Argentinos',
            'agrupacion': 'Provincia',
            'detalle': label,
            'total': value,
        })

    for label, value in zip(data.get('series', {}).get('extranjeros', {}).get('labels', []), data.get('series', {}).get('extranjeros', {}).get('values', [])):
        datos.append({
            'segmento': 'Extranjeros',
            'agrupacion': 'Continente',
            'detalle': label,
            'total': value,
        })

    return datos, 'Reporte Geográfico de Huéspedes'


def _reporte_importaciones():
    datos = db.ejecutar_query("""
        SELECT il.archivo_nombre, il.registros_importados,
               il.registros_duplicados, il.registros_error,
               il.fecha_importacion,
               u.username as usuario
        FROM importaciones_log il
        LEFT JOIN usuarios u ON il.usuario_id = u.id
        ORDER BY il.fecha_importacion DESC
    """, fetch=True)
    return [dict(d) for d in datos] if datos else [], "Historial de Importaciones"


def _reporte_auditoria():
    datos = db.ejecutar_query("""
        SELECT a.fecha, u.username, a.accion, a.detalle
        FROM auditoria a
        LEFT JOIN usuarios u ON a.usuario_id = u.id
        ORDER BY a.fecha DESC
        LIMIT 500
    """, fetch=True)
    return [dict(d) for d in datos] if datos else [], "Auditoría del Sistema"


def generar_excel(datos: list, titulo: str = "Datos") -> bytes:
    """Genera un archivo Excel en memoria y retorna los bytes."""
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = titulo[:31]

    if not datos:
        ws["A1"] = "Sin datos para exportar"
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1565C0", end_color="1565C0", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )
    alt_fill = PatternFill(start_color="E3F2FD", end_color="E3F2FD", fill_type="solid")

    column_names = {
        "id": "ID", "apellido_nombre": "Nombre y Apellido", "dni_pasaporte": "DNI/Pasaporte",
        "nacionalidad": "Nacionalidad", "procedencia": "Procedencia", "profesion": "Profesión",
        "edad": "Edad", "fecha_nacimiento": "Fecha Nac.", "fecha_entrada": "Entrada",
        "fecha_salida": "Salida", "hotel": "Hotel", "nombre": "Nombre", "hotel_nombre": "Hotel",
        "ciudad_localidad": "Ciudad", "direccion": "Dirección", "nro_orden": "Nro. Orden",
        "total": "Total", "username": "Usuario", "rol": "Rol", "ultimo_acceso": "Último Acceso",
        "fecha_registro": "Fecha Registro", "habitacion": "Habitación", "domicilio": "Domicilio",
        "destino": "Destino", "movilidad": "Movilidad", "telefono": "Teléfono",
        "categoria": "Categoría", "concepto": "Concepto", "archivo_nombre": "Archivo",
        "registros_importados": "Importados", "registros_duplicados": "Duplicados",
        "registros_error": "Errores", "fecha_importacion": "Fecha Importación",
        "usuario": "Usuario", "accion": "Acción", "detalle": "Detalle", "fecha": "Fecha",
    }

    # Título
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(datos[0]))
    title_cell = ws.cell(row=1, column=1, value=f"S.C.A.H. - {titulo}")
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="1565C0")
    title_cell.alignment = Alignment(horizontal="center")

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(datos[0]))
    ws.cell(row=2, column=1,
            value=f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}").font = Font(
        name="Calibri", size=9, italic=True, color="666666")

    columnas = list(datos[0].keys())
    for col_idx, col in enumerate(columnas, 1):
        cell = ws.cell(row=4, column=col_idx, value=column_names.get(col, col.title()))
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    for row_idx, fila in enumerate(datos, 5):
        for col_idx, col in enumerate(columnas, 1):
            valor = fila.get(col, "")
            if hasattr(valor, 'strftime'):
                valor = valor.strftime("%d/%m/%Y")
            elif valor is None:
                valor = ""
            cell = ws.cell(row=row_idx, column=col_idx, value=valor)
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)
            if (row_idx - 5) % 2 == 1:
                cell.fill = alt_fill

    for col_idx, col in enumerate(columnas, 1):
        max_len = len(column_names.get(col, col))
        for fila in datos[:100]:
            val = str(fila.get(col, ""))
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = min(max_len + 4, 40)

    total_row = len(datos) + 5
    ws.cell(row=total_row, column=1, value=f"Total de registros: {len(datos)}").font = Font(
        name="Calibri", size=10, bold=True)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def generar_pdf(datos: list, titulo: str = "Reporte", columnas_mostrar: list = None) -> bytes:
    """Genera un reporte PDF en memoria y retorna los bytes."""
    import io
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        leftMargin=15*mm, rightMargin=15*mm,
        topMargin=15*mm, bottomMargin=15*mm
    )

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="TituloReporte", parent=styles["Title"],
        fontSize=16, textColor=colors.HexColor("#1565C0"), spaceAfter=5))
    styles.add(ParagraphStyle(name="SubTitulo", parent=styles["Normal"],
        fontSize=9, textColor=colors.HexColor("#666666"), spaceAfter=15))
    styles.add(ParagraphStyle(name="CellText", parent=styles["Normal"],
        fontSize=7, leading=9))

    elements = []
    elements.append(Paragraph(f"S.C.A.H. - {titulo}", styles["TituloReporte"]))
    elements.append(Paragraph(
        f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Total: {len(datos)}",
        styles["SubTitulo"]))

    if not datos:
        elements.append(Paragraph("No hay datos para mostrar", styles["Normal"]))
        doc.build(elements)
        buffer.seek(0)
        return buffer.getvalue()

    column_names = {
        "apellido_nombre": "Nombre", "dni_pasaporte": "DNI", "nacionalidad": "Nacionalidad",
        "procedencia": "Procedencia", "profesion": "Profesión", "edad": "Edad",
        "fecha_entrada": "Entrada", "fecha_salida": "Salida", "hotel": "Hotel",
        "ciudad_localidad": "Ciudad", "habitacion": "Hab.", "domicilio": "Domicilio",
        "destino": "Destino", "movilidad": "Movilidad", "telefono": "Teléfono",
        "categoria": "Categoría", "concepto": "Concepto", "total": "Total",
        "nombre": "Nombre", "nro_orden": "Nro. Orden", "direccion": "Dirección",
        "total_huespedes": "Total Huéspedes", "activo": "Estado",
        "archivo_nombre": "Archivo", "registros_importados": "Importados",
        "username": "Usuario", "accion": "Acción", "detalle": "Detalle", "fecha": "Fecha",
    }

    if columnas_mostrar:
        cols = [c for c in columnas_mostrar if c in datos[0]]
    else:
        cols = list(datos[0].keys())
        if len(cols) > 1 and "id" in cols:
            cols.remove("id")

    headers = [column_names.get(c, c.title()) for c in cols]
    table_data = [headers]

    for fila in datos:
        row = []
        for c in cols:
            val = fila.get(c, "")
            if hasattr(val, 'strftime'):
                val = val.strftime("%d/%m/%Y")
            elif val is None:
                val = ""
            val = str(val)
            if len(val) > 30:
                val = val[:27] + "..."
            row.append(Paragraph(val, styles["CellText"]))
        table_data.append(row)

    available_width = landscape(A4)[0] - 30*mm
    col_widths = [available_width / len(cols)] * len(cols)

    table = Table(table_data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1565C0")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#E3F2FD")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))

    elements.append(table)
    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()
