"""
S.C.A.H. Web - Servicio de Importación Excel
Procesamiento de archivos Excel (formato V1 y V2 tabular).
"""

import sys
import os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..'))

from datetime import datetime
from openpyxl import load_workbook
from psycopg2.extras import execute_values
from database.connection import db
from config import (EXCEL_HOTEL_MAP, EXCEL_HUESPED_COLS, EXCEL_HUESPED_START_ROW,
                    EXCEL_V2_HUESPED_COLS, EXCEL_V2_HUESPED_START_ROW,
                    EXCEL_V2_HEADER_ALIASES)
from utils.validators import (validar_fecha, validar_edad, validar_telefono,
                               validar_habitacion, sanitizar_texto)
from utils.formatters import formato_fecha
from utils.logger import log_info, log_error, Auditoria

try:
    import xlrd
    XLRD_DISPONIBLE = True
except ImportError:
    XLRD_DISPONIBLE = False


class _XlrdCelda:
    def __init__(self, valor):
        self.value = valor


class _XlrdHoja:
    def __init__(self, sheet, book):
        self._sheet = sheet
        self._book = book
        self.title = sheet.name

    def __getitem__(self, ref: str):
        import re
        m = re.match(r'^([A-Za-z]+)(\d+)$', ref)
        if not m:
            return _XlrdCelda(None)

        col_letters = m.group(1).upper()
        row_num = int(m.group(2))

        col_idx = 0
        for ch in col_letters:
            col_idx = col_idx * 26 + (ord(ch) - ord('A') + 1)
        col_idx -= 1
        row_idx = row_num - 1

        if row_idx < 0 or row_idx >= self._sheet.nrows:
            return _XlrdCelda(None)
        if col_idx < 0 or col_idx >= self._sheet.ncols:
            return _XlrdCelda(None)

        cell = self._sheet.cell(row_idx, col_idx)
        valor = cell.value

        if cell.ctype == xlrd.XL_CELL_DATE and valor:
            try:
                dt_tuple = xlrd.xldate_as_tuple(valor, self._book.datemode)
                from datetime import datetime as _dt
                valor = _dt(*dt_tuple)
            except Exception:
                pass

        if isinstance(valor, str) and valor.strip() == '':
            valor = None

        return _XlrdCelda(valor)


class _XlrdLibro:
    def __init__(self, filepath):
        self._book = xlrd.open_workbook(filepath)
        self.sheetnames = self._book.sheet_names()

    def __getitem__(self, nombre):
        sheet = self._book.sheet_by_name(nombre)
        return _XlrdHoja(sheet, self._book)

    def close(self):
        self._book.release_resources()


def _abrir_workbook(filepath: str):
    ext = os.path.splitext(filepath)[1].lower()
    if ext == '.xls':
        if not XLRD_DISPONIBLE:
            raise ImportError(
                "Se necesita la librería 'xlrd' para leer archivos .xls. "
                "Instálala con: pip install xlrd"
            )
        return _XlrdLibro(filepath)
    return load_workbook(filepath, read_only=True, data_only=True)


# ── Lectura de celdas ────────────────────────────────────────

def _valor_celda(hoja, columna: str, fila: int):
    try:
        celda = hoja[f"{columna}{fila}"]
        valor = celda.value
        if isinstance(valor, str) and valor.strip() == "":
            return None
        return valor
    except Exception:
        return None


def _fila_tiene_datos(hoja, fila: int, cols: dict) -> bool:
    for col in cols.values():
        valor = _valor_celda(hoja, col, fila)
        if valor is not None and str(valor).strip() != "":
            return True
    return False


def _obtener_max_fila(hoja) -> int:
    if hasattr(hoja, 'max_row') and hoja.max_row:
        return hoja.max_row
    if hasattr(hoja, '_sheet') and hasattr(hoja._sheet, 'nrows'):
        return hoja._sheet.nrows
    return 10000


# ── Formato V1 (hotel+huéspedes por hoja) ────────────────────

_ENCABEZADOS = {
    "apellido", "nombre", "apellido y nombre", "apellido_nombre",
    "dni", "pasaporte", "dni/pasaporte", "documento",
    "nacionalidad", "procedencia", "profesion", "profesión",
    "edad", "entrada", "salida", "nacimiento",
    "fecha", "hotel", "nro", "orden", "dirección", "direccion",
    "ciudad", "localidad", "fec. nacimiento", "fec. entrada",
    "fec. salida", "fecha nacimiento", "fecha entrada", "fecha salida",
    "habitacion", "habitación",
}


def _detectar_fila_inicio_v1(hoja) -> int:
    columnas_a_verificar = [
        EXCEL_HUESPED_COLS["apellido_nombre"],
        EXCEL_HUESPED_COLS["dni_pasaporte"],
        EXCEL_HUESPED_COLS["nacionalidad"],
    ]
    for fila in range(1, 50):
        for col in columnas_a_verificar:
            valor = _valor_celda(hoja, col, fila)
            if valor is not None:
                texto = str(valor).strip().lower()
                if texto and texto not in _ENCABEZADOS and len(texto) > 1:
                    return fila
    return EXCEL_HUESPED_START_ROW


def _leer_datos_hotel(hoja, nombre_archivo: str, hoja_nombre: str) -> dict:
    try:
        nombre = _valor_celda(hoja, EXCEL_HOTEL_MAP["nombre"]["col"],
                              EXCEL_HOTEL_MAP["nombre"]["row"])
        nro_orden = _valor_celda(hoja, EXCEL_HOTEL_MAP["nro_orden"]["col"],
                                 EXCEL_HOTEL_MAP["nro_orden"]["row"])
        direccion = _valor_celda(hoja, EXCEL_HOTEL_MAP["direccion"]["col"],
                                 EXCEL_HOTEL_MAP["direccion"]["row"])
        ciudad = _valor_celda(hoja, EXCEL_HOTEL_MAP["ciudad_localidad"]["col"],
                              EXCEL_HOTEL_MAP["ciudad_localidad"]["row"])

        if not nombre or str(nombre).strip() == "":
            for r in range(1, 10):
                val = _valor_celda(hoja, "A", r)
                if val and str(val).strip() and len(str(val).strip()) > 2:
                    texto = str(val).strip().lower()
                    if texto not in ("hotel", "nombre", "nro", "orden", "dirección",
                                     "direccion", "ciudad", "localidad"):
                        nombre = val
                        break

        if not nombre or str(nombre).strip() == "":
            nombre = hoja_nombre

        return {
            "archivo_hoja": f"{nombre_archivo} / {hoja_nombre}",
            "nombre": sanitizar_texto(str(nombre)) if nombre else f"Hotel ({hoja_nombre})",
            "nro_orden": sanitizar_texto(str(nro_orden)) if nro_orden else "",
            "direccion": sanitizar_texto(str(direccion)) if direccion else "",
            "ciudad": sanitizar_texto(str(ciudad)) if ciudad else "",
            "hoja": hoja_nombre,
        }
    except Exception as e:
        log_error(f"Error leyendo datos de hotel en {hoja_nombre}", e)
        return {
            "archivo_hoja": f"{nombre_archivo} / {hoja_nombre}",
            "nombre": f"Hotel ({hoja_nombre})",
            "nro_orden": "", "direccion": "", "ciudad": "",
            "hoja": hoja_nombre,
        }


def _leer_huespedes_v1(hoja, hotel_data: dict) -> list:
    huespedes = []
    fila = _detectar_fila_inicio_v1(hoja)
    max_fila = _obtener_max_fila(hoja)
    filas_vacias = 0

    while fila <= max_fila + 15:
        if not _fila_tiene_datos(hoja, fila, EXCEL_HUESPED_COLS):
            filas_vacias += 1
            if filas_vacias >= 15:
                break
            fila += 1
            continue
        filas_vacias = 0

        apellido = _valor_celda(hoja, EXCEL_HUESPED_COLS["apellido_nombre"], fila)
        if not apellido or str(apellido).strip() == "":
            dni_val = _valor_celda(hoja, EXCEL_HUESPED_COLS["dni_pasaporte"], fila)
            if not dni_val or str(dni_val).strip() == "":
                fila += 1
                continue

        try:
            fn_raw = _valor_celda(hoja, EXCEL_HUESPED_COLS["fecha_nacimiento"], fila)
            _, _, fecha_nac = validar_fecha(fn_raw, permite_vacio=True)
            fe_raw = _valor_celda(hoja, EXCEL_HUESPED_COLS["fecha_entrada"], fila)
            _, _, fecha_entrada = validar_fecha(fe_raw, permite_vacio=True)
            fs_raw = _valor_celda(hoja, EXCEL_HUESPED_COLS["fecha_salida"], fila)
            _, _, fecha_salida = validar_fecha(fs_raw, permite_vacio=True)
            edad_raw = _valor_celda(hoja, EXCEL_HUESPED_COLS["edad"], fila)
            _, _, edad = validar_edad(edad_raw)

            huesped = {
                "hotel_nombre": hotel_data.get("nombre", "") if hotel_data else "",
                "hotel_data": hotel_data,
                "nacionalidad": sanitizar_texto(
                    str(_valor_celda(hoja, EXCEL_HUESPED_COLS["nacionalidad"], fila) or "")),
                "procedencia": sanitizar_texto(
                    str(_valor_celda(hoja, EXCEL_HUESPED_COLS["procedencia"], fila) or "")),
                "apellido_nombre": sanitizar_texto(str(apellido)),
                "dni_pasaporte": sanitizar_texto(
                    str(_valor_celda(hoja, EXCEL_HUESPED_COLS["dni_pasaporte"], fila) or "")),
                "fecha_nacimiento": fecha_nac,
                "edad": edad,
                "profesion": sanitizar_texto(
                    str(_valor_celda(hoja, EXCEL_HUESPED_COLS["profesion"], fila) or "")),
                "fecha_entrada": fecha_entrada,
                "fecha_salida": fecha_salida,
            }
            huespedes.append(huesped)
        except Exception as e:
            log_error(f"Error leyendo huésped en fila {fila}", e)

        fila += 1
    return huespedes


# ── Formato V2 (tabular con encabezados) ─────────────────────

def _detectar_encabezados(hoja) -> dict:
    mapeo = {}
    for col_idx in range(26):
        letra = chr(ord('A') + col_idx)
        try:
            celda = hoja[f"{letra}1"]
            valor = celda.value
            if valor is None:
                continue
            texto = str(valor).strip().lower()
            texto_limpio = (texto.replace("á", "a").replace("é", "e")
                           .replace("í", "i").replace("ó", "o")
                           .replace("ú", "u").replace("ñ", "n"))
            if texto in EXCEL_V2_HEADER_ALIASES:
                campo = EXCEL_V2_HEADER_ALIASES[texto]
                if campo not in mapeo:
                    mapeo[campo] = letra
            elif texto_limpio in EXCEL_V2_HEADER_ALIASES:
                campo = EXCEL_V2_HEADER_ALIASES[texto_limpio]
                if campo not in mapeo:
                    mapeo[campo] = letra
        except Exception:
            continue
    return mapeo


def _detectar_fila_inicio_v2(hoja, mapeo: dict) -> int:
    columna_nombre = mapeo.get("apellido_nombre", "C")
    for fila in range(1, 20):
        valor = _valor_celda(hoja, columna_nombre, fila)
        if valor is not None:
            texto = str(valor).strip().lower()
            if texto and texto not in _ENCABEZADOS and len(texto) > 1:
                return fila
    return EXCEL_V2_HUESPED_START_ROW


def _leer_huespedes_v2(hoja, mapeo: dict, nombre_archivo: str, hoja_nombre: str) -> list:
    huespedes = []
    fila = _detectar_fila_inicio_v2(hoja, mapeo)
    max_fila = _obtener_max_fila(hoja)
    filas_vacias = 0

    while fila <= max_fila + 15:
        if not _fila_tiene_datos(hoja, fila, mapeo):
            filas_vacias += 1
            if filas_vacias >= 15:
                break
            fila += 1
            continue
        filas_vacias = 0

        col_nombre = mapeo.get("apellido_nombre", "C")
        apellido = _valor_celda(hoja, col_nombre, fila)
        if not apellido or str(apellido).strip() == "":
            col_dni = mapeo.get("dni_pasaporte", "J")
            dni_val = _valor_celda(hoja, col_dni, fila)
            if not dni_val or str(dni_val).strip() == "":
                fila += 1
                continue

        try:
            fe_raw = _valor_celda(hoja, mapeo.get("fecha_entrada", "A"), fila) if "fecha_entrada" in mapeo else None
            _, _, fecha_entrada = validar_fecha(fe_raw, permite_vacio=True)
            fs_raw = _valor_celda(hoja, mapeo.get("fecha_salida", "K"), fila) if "fecha_salida" in mapeo else None
            _, _, fecha_salida = validar_fecha(fs_raw, permite_vacio=True)
            edad_raw = _valor_celda(hoja, mapeo.get("edad", "D"), fila) if "edad" in mapeo else None
            _, _, edad = validar_edad(edad_raw)
            hab_raw = _valor_celda(hoja, mapeo.get("habitacion", "B"), fila) if "habitacion" in mapeo else None
            _, _, habitacion = validar_habitacion(hab_raw)
            tel_raw = _valor_celda(hoja, mapeo.get("telefono", "M"), fila) if "telefono" in mapeo else None
            _, _, telefono = validar_telefono(tel_raw)

            dni_raw = _valor_celda(hoja, mapeo.get("dni_pasaporte", "J"), fila) if "dni_pasaporte" in mapeo else None
            dni_str = sanitizar_texto(str(dni_raw)) if dni_raw else ""
            if dni_str.endswith('.0'):
                dni_str = dni_str[:-2]

            huesped = {
                "archivo": f"{nombre_archivo}/{hoja_nombre}",
                "fecha_entrada": fecha_entrada,
                "habitacion": habitacion,
                "apellido_nombre": sanitizar_texto(str(apellido)) if apellido else "",
                "edad": edad,
                "nacionalidad": sanitizar_texto(
                    str(_valor_celda(hoja, mapeo.get("nacionalidad", "E"), fila) or ""))
                    if "nacionalidad" in mapeo else "",
                "profesion": sanitizar_texto(
                    str(_valor_celda(hoja, mapeo.get("profesion", "F"), fila) or ""))
                    if "profesion" in mapeo else "",
                "procedencia": sanitizar_texto(
                    str(_valor_celda(hoja, mapeo.get("procedencia", "G"), fila) or ""))
                    if "procedencia" in mapeo else "",
                "domicilio": sanitizar_texto(
                    str(_valor_celda(hoja, mapeo.get("domicilio", "H"), fila) or ""))
                    if "domicilio" in mapeo else "",
                "destino": sanitizar_texto(
                    str(_valor_celda(hoja, mapeo.get("destino", "I"), fila) or ""))
                    if "destino" in mapeo else "",
                "dni_pasaporte": dni_str,
                "fecha_salida": fecha_salida,
                "movilidad": sanitizar_texto(
                    str(_valor_celda(hoja, mapeo.get("movilidad", "L"), fila) or ""))
                    if "movilidad" in mapeo else "",
                "telefono": telefono,
            }
            huespedes.append(huesped)
        except Exception as e:
            log_error(f"Error leyendo huésped fila {fila} de {nombre_archivo}", e)

        fila += 1
    return huespedes


# ── API pública ───────────────────────────────────────────────

def procesar_archivos_v1(rutas_archivos: list) -> dict:
    """
    Procesa archivos Excel en formato V1 (hotel por hoja).
    Retorna {hoteles: [...], huespedes: [...], errores: [...]}
    """
    hoteles = []
    huespedes = []
    errores = []

    for archivo in rutas_archivos:
        try:
            wb = _abrir_workbook(archivo)
            nombre_archivo = os.path.basename(archivo)

            for hoja_nombre in wb.sheetnames:
                hoja = wb[hoja_nombre]
                hotel_data = _leer_datos_hotel(hoja, nombre_archivo, hoja_nombre)
                if hotel_data:
                    hoteles.append(hotel_data)
                huesps = _leer_huespedes_v1(hoja, hotel_data)
                huespedes.extend(huesps)

            wb.close()
        except Exception as e:
            error_msg = f"Error en '{os.path.basename(archivo)}': {str(e)}"
            errores.append(error_msg)
            log_error(error_msg, e)

    return {"hoteles": hoteles, "huespedes": huespedes, "errores": errores}


def procesar_archivos_v2(rutas_archivos: list) -> dict:
    """
    Procesa archivos Excel en formato V2 (tabular con encabezados).
    Retorna {huespedes: [...], errores: [...], columnas_detectadas: {...}}
    """
    huespedes = []
    errores = []
    mapeo_columnas = {}

    for archivo in rutas_archivos:
        try:
            wb = _abrir_workbook(archivo)
            nombre_archivo = os.path.basename(archivo)

            for hoja_nombre in wb.sheetnames:
                hoja = wb[hoja_nombre]
                mapeo_detectado = _detectar_encabezados(hoja)
                if not mapeo_detectado or len(mapeo_detectado) < 3:
                    mapeo_detectado = dict(EXCEL_V2_HUESPED_COLS)
                mapeo_columnas = mapeo_detectado
                huesps = _leer_huespedes_v2(hoja, mapeo_detectado, nombre_archivo, hoja_nombre)
                huespedes.extend(huesps)

            wb.close()
        except Exception as e:
            error_msg = f"Error en '{os.path.basename(archivo)}': {str(e)}"
            errores.append(error_msg)
            log_error(error_msg, e)

    return {
        "huespedes": huespedes,
        "errores": errores,
        "columnas_detectadas": mapeo_columnas,
    }


def importar_datos_v1(huespedes: list, usuario_id: int) -> dict:
    """
    Importa huéspedes del formato V1 a la base de datos.
    Crea hoteles automáticamente si no existen.
    Retorna {importados, duplicados, errores}
    """
    importados = 0
    errores_count = 0
    duplicados = 0

    conn = db.obtener_conexion()
    if not conn:
        return {"importados": 0, "duplicados": 0, "errores": 1,
                "mensaje": "Error de conexión a la base de datos"}

    descartar = False
    try:
        cursor = conn.cursor()

        # 1) Resolver los hoteles una sola vez cada uno (no una vez por fila).
        hoteles_cache = {}
        for huesped in huespedes:
            hotel_data = huesped.get("hotel_data", {}) or {}
            hotel_key = hotel_data.get("nombre", "")
            if hotel_key not in hoteles_cache:
                hoteles_cache[hotel_key] = _obtener_o_crear_hotel(
                    cursor, hotel_data, usuario_id)

        # 2) Traer de una vez lo ya cargado, para comparar en memoria.
        claves_existentes = _cargar_claves_existentes(cursor, hoteles_cache.values())

        # 3) Preparar las filas a insertar.
        filas = []
        for huesped in huespedes:
            try:
                hotel_data = huesped.get("hotel_data", {}) or {}
                hotel_id = hoteles_cache.get(hotel_data.get("nombre", ""))

                clave = _clave_duplicado(hotel_id, huesped)
                if clave is not None and clave in claves_existentes:
                    duplicados += 1
                    continue
                if clave is not None:
                    # Así se descartan también los repetidos dentro del archivo.
                    claves_existentes.add(clave)

                filas.append((
                    hotel_id, huesped["nacionalidad"], huesped["procedencia"],
                    huesped["apellido_nombre"], huesped["dni_pasaporte"],
                    huesped.get("fecha_nacimiento"), huesped.get("edad"),
                    huesped["profesion"], huesped.get("fecha_entrada"),
                    huesped.get("fecha_salida"), usuario_id
                ))
            except Exception as e:
                errores_count += 1
                log_error(f"Error preparando huésped: {huesped.get('apellido_nombre', '?')}", e)

        # 4) Insertar agrupado.
        insertadas, errores_lote = _insertar_en_lotes(cursor, """
            INSERT INTO huespedes (
                hotel_id, nacionalidad, procedencia, apellido_nombre,
                dni_pasaporte, fecha_nacimiento, edad, profesion,
                fecha_entrada, fecha_salida, usuario_carga_id, origen_carga
            ) VALUES %s
        """, [fila + ('excel',) for fila in filas])
        importados = insertadas
        errores_count += errores_lote

        _registrar_log_importacion(cursor, "excel_v1", importados, errores_count,
                                    duplicados, usuario_id)
        conn.commit()
        cursor.close()

    except Exception as e:
        descartar = _revertir_conexion(conn)
        log_error("Error general en importación V1", e)
        return {"importados": 0, "duplicados": duplicados,
                "errores": errores_count + 1, "mensaje": str(e)}
    finally:
        # Sin este finally, una importación fallida se quedaba con la conexión
        # para siempre: unas pocas bastaban para agotar el pool y dejar toda la
        # aplicación bloqueada hasta reiniciarla.
        db.liberar_conexion(conn, descartar=descartar)

    _invalidar_caches()
    _registrar_auditoria(usuario_id, "importar_excel",
                         f"V1 - Importados: {importados}, Errores: {errores_count}, "
                         f"Duplicados: {duplicados}")

    return {"importados": importados, "duplicados": duplicados,
            "errores": errores_count,
            "mensaje": f"Importados: {importados}, Duplicados: {duplicados}, Errores: {errores_count}"}


def importar_datos_v2(huespedes: list, hotel_nombre: str, usuario_id: int,
                      nuevo_hotel: bool = False, hotel_extra: dict = None) -> dict:
    """
    Importa huéspedes del formato V2 a la base de datos.
    Requiere selección explícita de hotel.
    Retorna {importados, duplicados, errores}
    """
    importados = 0
    errores_count = 0
    duplicados = 0

    conn = db.obtener_conexion()
    if not conn:
        return {"importados": 0, "duplicados": 0, "errores": 1,
                "mensaje": "Error de conexión a la base de datos"}

    descartar = False
    try:
        cursor = conn.cursor()

        if nuevo_hotel:
            extra = hotel_extra or {}
            cursor.execute("""
                INSERT INTO hoteles (nombre, nro_orden, direccion, ciudad_localidad, usuario_registro_id)
                VALUES (%s,%s,%s,%s,%s) RETURNING id
            """, (
                sanitizar_texto(hotel_nombre),
                sanitizar_texto(extra.get("nro_orden", "")),
                sanitizar_texto(extra.get("direccion", "")),
                sanitizar_texto(extra.get("ciudad", "")),
                usuario_id
            ))
            hotel_id = cursor.fetchone()[0]
        else:
            cursor.execute("SELECT id FROM hoteles WHERE LOWER(nombre) = LOWER(%s)",
                           (hotel_nombre,))
            resultado = cursor.fetchone()
            if resultado:
                hotel_id = resultado[0]
            else:
                cursor.execute("""
                    INSERT INTO hoteles (nombre, usuario_registro_id)
                    VALUES (%s,%s) RETURNING id
                """, (sanitizar_texto(hotel_nombre), usuario_id))
                hotel_id = cursor.fetchone()[0]

        # Una sola consulta para saber qué hay ya cargado en este hotel.
        claves_existentes = _cargar_claves_existentes(cursor, [hotel_id])

        filas = []
        for huesped in huespedes:
            try:
                clave = _clave_duplicado(hotel_id, huesped)
                if clave is not None and clave in claves_existentes:
                    duplicados += 1
                    continue
                if clave is not None:
                    claves_existentes.add(clave)

                filas.append((
                    hotel_id, huesped.get("nacionalidad", ""),
                    huesped.get("procedencia", ""),
                    huesped.get("apellido_nombre", ""),
                    huesped.get("dni_pasaporte", ""),
                    huesped.get("edad"), huesped.get("profesion", ""),
                    huesped.get("fecha_entrada"), huesped.get("fecha_salida"),
                    huesped.get("habitacion", ""), huesped.get("domicilio", ""),
                    huesped.get("destino", ""), huesped.get("movilidad", ""),
                    huesped.get("telefono", ""), usuario_id, 'excel_v2'
                ))
            except Exception as e:
                errores_count += 1
                log_error(f"Error preparando: {huesped.get('apellido_nombre', '?')}", e)

        insertadas, errores_lote = _insertar_en_lotes(cursor, """
            INSERT INTO huespedes (
                hotel_id, nacionalidad, procedencia, apellido_nombre,
                dni_pasaporte, edad, profesion,
                fecha_entrada, fecha_salida,
                habitacion, domicilio, destino, movilidad, telefono,
                usuario_carga_id, origen_carga
            ) VALUES %s
        """, filas)
        importados = insertadas
        errores_count += errores_lote

        _registrar_log_importacion(cursor, "excel_v2", importados, errores_count,
                                    duplicados, usuario_id)
        conn.commit()
        cursor.close()

    except Exception as e:
        descartar = _revertir_conexion(conn)
        log_error("Error general en importación V2", e)
        return {"importados": 0, "duplicados": duplicados,
                "errores": errores_count + 1, "mensaje": str(e)}
    finally:
        db.liberar_conexion(conn, descartar=descartar)

    _invalidar_caches()
    _registrar_auditoria(usuario_id, "importar_excel_v2",
                         f"Tabular - Importados: {importados}, Errores: {errores_count}, "
                         f"Duplicados: {duplicados}")

    return {"importados": importados, "duplicados": duplicados,
            "errores": errores_count,
            "mensaje": f"Importados: {importados}, Duplicados: {duplicados}, Errores: {errores_count}"}


def obtener_hoteles_activos() -> list:
    """Retorna lista de hoteles activos para selector."""
    resultado = db.ejecutar_query(
        "SELECT id, nombre, ciudad_localidad FROM hoteles WHERE activo = TRUE ORDER BY nombre",
        fetch=True
    )
    return [dict(r) for r in resultado] if resultado else []


# ── Funciones auxiliares internas ─────────────────────────────

def _obtener_o_crear_hotel(cursor, hotel_data: dict, usuario_id: int) -> int:
    nombre = hotel_data.get("nombre", "Sin nombre")
    cursor.execute("SELECT id FROM hoteles WHERE LOWER(nombre) = LOWER(%s)", (nombre,))
    resultado = cursor.fetchone()
    if resultado:
        return resultado[0]
    cursor.execute("""
        INSERT INTO hoteles (nombre, nro_orden, direccion, ciudad_localidad, usuario_registro_id)
        VALUES (%s,%s,%s,%s,%s) RETURNING id
    """, (nombre, hotel_data.get("nro_orden", ""), hotel_data.get("direccion", ""),
          hotel_data.get("ciudad", ""), usuario_id))
    return cursor.fetchone()[0]


# ── Importación por lotes ────────────────────────────────────
# Tamaño de página para los INSERT agrupados. 500 filas por sentencia mantiene
# el tamaño del mensaje razonable y reduce los viajes de red ~500 veces.
_BATCH_PAGE_SIZE = 500


def _clave_duplicado(hotel_id: int, huesped: dict):
    """Clave de duplicidad (hotel, documento, fecha de entrada).

    Retorna None cuando faltan documento o fecha: en ese caso el registro nunca
    se considera duplicado, igual que hacía la verificación fila a fila.
    """
    dni = huesped.get("dni_pasaporte", "")
    fecha_entrada = huesped.get("fecha_entrada")
    if not dni or not fecha_entrada:
        return None
    return (hotel_id, dni, fecha_entrada)


def _cargar_claves_existentes(cursor, hotel_ids) -> set:
    """Trae en UNA consulta los registros ya cargados de los hoteles implicados.

    Sustituye al SELECT por fila: con un Excel de 2.000 filas eran 2.000 viajes
    de ida y vuelta a la base de datos, que contra una base remota bastaban para
    superar el timeout del worker y dejar la aplicación sin capacidad.
    """
    ids = [h for h in set(hotel_ids) if h]
    if not ids:
        return set()
    cursor.execute("""
        SELECT hotel_id, dni_pasaporte, fecha_entrada
        FROM huespedes
        WHERE hotel_id = ANY(%s)
          AND dni_pasaporte IS NOT NULL AND dni_pasaporte <> ''
          AND fecha_entrada IS NOT NULL
    """, (ids,))
    return {(fila[0], fila[1], fila[2]) for fila in cursor.fetchall()}


def _insertar_en_lotes(cursor, sql: str, filas: list) -> tuple[int, int]:
    """Inserta las filas agrupadas. Retorna (insertadas, con_error).

    Si un lote falla se reintenta fila a fila sobre un SAVEPOINT, de modo que un
    único registro defectuoso no invalide el resto de la importación (que es lo
    que hacía la versión fila a fila).
    """
    insertadas = 0
    errores = 0

    for inicio in range(0, len(filas), _BATCH_PAGE_SIZE):
        lote = filas[inicio:inicio + _BATCH_PAGE_SIZE]
        cursor.execute("SAVEPOINT lote_import")
        try:
            execute_values(cursor, sql, lote, page_size=len(lote))
            cursor.execute("RELEASE SAVEPOINT lote_import")
            insertadas += len(lote)
        except Exception as e:
            cursor.execute("ROLLBACK TO SAVEPOINT lote_import")
            log_error(f"Lote de importación con error, se reintenta fila a fila "
                      f"({len(lote)} registros)", e)
            for fila in lote:
                cursor.execute("SAVEPOINT fila_import")
                try:
                    execute_values(cursor, sql, [fila], page_size=1)
                    cursor.execute("RELEASE SAVEPOINT fila_import")
                    insertadas += 1
                except Exception as e_fila:
                    cursor.execute("ROLLBACK TO SAVEPOINT fila_import")
                    errores += 1
                    log_error("Error importando registro", e_fila)

    return insertadas, errores


def _registrar_log_importacion(cursor, tipo: str, importados: int,
                                errores: int, duplicados: int, usuario_id: int):
    cursor.execute("""
        INSERT INTO importaciones_log (
            archivo_nombre, fecha_importacion, usuario_id,
            registros_importados, registros_error, registros_duplicados, estado
        ) VALUES (%s, %s, %s, %s, %s, %s, %s)
    """, (
        f"upload_{tipo}_{datetime.now().strftime('%Y%m%d_%H%M%S')}",
        datetime.now(), usuario_id, importados, errores, duplicados,
        "completado" if errores == 0 else "parcial"
    ))


def _registrar_auditoria(usuario_id: int, accion: str, detalle: str):
    conn = None
    try:
        conn = db.obtener_conexion()
        if conn:
            auditoria = Auditoria(conn)
            auditoria.registrar(usuario_id, accion, "huespedes", detalle=detalle)
    except Exception:
        pass
    finally:
        # En finally: si registrar() lanzaba, la conexión no volvía al pool.
        if conn is not None:
            db.liberar_conexion(conn)


def _revertir_conexion(conn) -> bool:
    """Deshace la transacción tras un error. True si la conexión quedó inservible."""
    try:
        conn.rollback()
        return False
    except Exception:
        return True


def _invalidar_caches():
    """Descarta los totales cacheados de los listados tras una importación."""
    try:
        from web.services.guest_service import invalidar_cache_conteos
        invalidar_cache_conteos()
    except Exception:
        pass
